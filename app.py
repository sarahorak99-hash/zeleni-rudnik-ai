"""Zeleni rudnik – verzija 2: karta dvorišta + prepoznavanje otpada sa slike.

Pokretanje (u mapi s ovim datotekama):
    py -3.13 -m streamlit run app.py

Mapa mora sadržavati:
    app.py
    reciklazna_dvorista_ociscena.csv
    Zeleni_rudnik_tablica_predmeta.xlsx   (samo za prepoznavanje otpada)

Za prepoznavanje trebaju još: torch, transformers, pillow, openpyxl.
Ako ih nema, stranica karte i dalje radi.
"""
from datetime import datetime
from math import asin, cos, radians, sin, sqrt
from zoneinfo import ZoneInfo

import folium
import pandas as pd
import streamlit as st
from streamlit_folium import st_folium

CSV_PATH = "reciklazna_dvorista_ociscena.csv"
TABLE_PATH = "Zeleni_rudnik_tablica_predmeta.xlsx"
MODEL_NAME = "openai/clip-vit-base-patch32"
ZAGREB_CENTER = (45.8131, 15.9775)
CONFIDENCE_THRESHOLD = 0.60  # ispod ovoga aplikacija ne tvrdi rezultat
TOP_N = 3

# Predmeti čiji se sadržaj ne vidi na slici: tražimo potvrdu korisnika.
CONTENT_DEPENDENT = ["ulje", "lijekov", "pesticid", "zaštitu bilja", "aerosol", "boje", "otapal"]

# Ključna riječ u nazivu predmeta -> stupac u CSV-u dvorišta koji to dvorište mora imati.
YARD_FLAG_BY_KEYWORD = {
    "lijekov": "lijekovi",
    "tekstil": "tekstil",
    "odjeć": "tekstil",
    "akumulator": "akumulatori",
    "fluorescent": "fluorescentne_cijevi",
    "namještaj": "glomazni_otpad",
    "madrac": "glomazni_otpad",
}

TEXT = {
    "hr": {
        "title": "Zeleni rudnik",
        "subtitle": "Kamo s otpadom u Zagrebu? Pronađi najbliže reciklažno dvorište.",
        "page": "Stranica",
        "nav_map": "Karta dvorišta",
        "nav_scan": "Prepoznaj otpad",
        "filters": "Filteri",
        "type": "Vrsta dvorišta",
        "type_opts": {"stalno": "Stalno", "mobilno": "Mobilno"},
        "open_now": "Samo otvoreno sada",
        "extra": "Prima i",
        "extra_opts": {
            "lijekovi": "Stare lijekove",
            "tekstil": "Tekstil i odjeću",
            "akumulatori": "Akumulatore",
            "fluorescentne_cijevi": "Fluorescentne cijevi",
            "glomazni_otpad": "Glomazni otpad",
        },
        "click": "Klikni na kartu da označiš svoju lokaciju.",
        "nearest": "Najbliže dvorište",
        "distance": "Udaljenost (zračno)",
        "route": "Ruta (OpenStreetMap)",
        "open": "Otvoreno",
        "closed": "Zatvoreno",
        "hours": "Radno vrijeme",
        "phone": "Telefon",
        "none": "Nijedno dvorište ne odgovara filterima.",
        "you": "Tvoja lokacija",
        "note": "Napomena: praznici nisu uračunati, a raspored mobilnih dvorišta provjeri na cistoca.hr.",
        "scan_title": "Prepoznaj otpad",
        "scan_intro": "Slikaj predmet koji želiš baciti, a mi ćemo ti reći kamo s njim.",
        "src_cam": "Snimi kamerom",
        "src_upl": "Učitaj sliku",
        "take": "Slikaj predmet",
        "upload": "Odaberi sliku (jpg, png, webp)",
        "analysing": "Analiziram sliku...",
        "likely": "Najvjerojatnije je",
        "unsure": "Nisam sigurna što je na slici. Odaberi što je najbliže ili probaj ponovno s drugim kutom.",
        "pick": "Što je od ovoga točno?",
        "none_of": "Ništa od ovoga",
        "none_msg": "U redu. Probaj slikati predmet bliže, na jednoj podlozi i uz dobro svjetlo.",
        "category": "Kategorija",
        "where": "Gdje zbrinuti",
        "remark": "Napomena",
        "sensitive": "Ovo ne možemo pouzdano prepoznati sa slike jer ovisi o sadržaju.",
        "confirm": "Potvrđujem da je to točno ono što je u ambalaži",
        "no_pos": "Za najbliže dvorište označi svoju lokaciju na stranici Karta dvorišta.",
        "unavailable": "Prepoznavanje otpada nije dostupno u ovoj verziji (nedostaju biblioteke ili tablica predmeta).",
        "no_table": "Nije pronađena tablica predmeta (Zeleni_rudnik_tablica_predmeta.xlsx).",
        "disclaimer": "Prijedlog je automatski i može biti pogrešan. Ako nisi sigurna, provjeri upute na cistoca.hr.",
    },
    "en": {
        "title": "Zeleni rudnik",
        "subtitle": "Where to take your waste in Zagreb? Find the nearest recycling yard.",
        "page": "Page",
        "nav_map": "Yard map",
        "nav_scan": "Identify waste",
        "filters": "Filters",
        "type": "Yard type",
        "type_opts": {"stalno": "Permanent", "mobilno": "Mobile"},
        "open_now": "Open now only",
        "extra": "Also accepts",
        "extra_opts": {
            "lijekovi": "Old medicines",
            "tekstil": "Textiles and clothes",
            "akumulatori": "Car batteries",
            "fluorescentne_cijevi": "Fluorescent tubes",
            "glomazni_otpad": "Bulky waste",
        },
        "click": "Click on the map to mark your location.",
        "nearest": "Nearest yard",
        "distance": "Distance (straight line)",
        "route": "Route (OpenStreetMap)",
        "open": "Open",
        "closed": "Closed",
        "hours": "Opening hours",
        "phone": "Phone",
        "none": "No yard matches the filters.",
        "you": "Your location",
        "note": "Note: public holidays are not included; check the mobile yard schedule at cistoca.hr.",
        "scan_title": "Identify waste",
        "scan_intro": "Take a photo of the item you want to throw away and we will tell you where it goes.",
        "src_cam": "Use camera",
        "src_upl": "Upload image",
        "take": "Take a photo of the item",
        "upload": "Choose an image (jpg, png, webp)",
        "analysing": "Analysing the image...",
        "likely": "Most likely it is",
        "unsure": "I am not sure what is in the picture. Pick the closest option or try again from another angle.",
        "pick": "Which of these is correct?",
        "none_of": "None of these",
        "none_msg": "OK. Try photographing the item closer, on a plain surface and in good light.",
        "category": "Category",
        "where": "Where to dispose",
        "remark": "Note",
        "sensitive": "We cannot reliably tell this from a photo because it depends on the contents.",
        "confirm": "I confirm this is exactly what is in the packaging",
        "no_pos": "To find the nearest yard, mark your location on the Yard map page.",
        "unavailable": "Waste identification is not available in this version (missing libraries or item table).",
        "no_table": "Item table not found (Zeleni_rudnik_tablica_predmeta.xlsx).",
        "disclaimer": "The suggestion is automatic and may be wrong. If unsure, check the instructions at cistoca.hr.",
    },
}


# ---------------------------------------------------------------- pomoćne funkcije
@st.cache_data
def load_yards() -> pd.DataFrame:
    df = pd.read_csv(CSV_PATH, encoding="utf-8", dtype=str, keep_default_na=False)
    df["lat"] = df["lat"].astype(float)
    df["lon"] = df["lon"].astype(float)
    for col in ["lijekovi", "tekstil", "akumulatori", "fluorescentne_cijevi", "glomazni_otpad"]:
        df[col] = df[col].str.lower().eq("true")
    return df


def to_minutes(hhmm: str):
    if not hhmm:
        return None
    h, m = hhmm.split(":")
    return int(h) * 60 + int(m)


def is_open(row, now: datetime) -> bool:
    """Pon-sub po rasporedu pon_sub_*, nedjelja po ned_* (ako postoji)."""
    if now.weekday() == 6:
        start, end = to_minutes(row["ned_od"]), to_minutes(row["ned_do"])
    else:
        start, end = to_minutes(row["pon_sub_od"]), to_minutes(row["pon_sub_do"])
    if start is None or end is None:
        return False
    current = now.hour * 60 + now.minute
    return start <= current < end


def haversine_km(lat1, lon1, lat2, lon2) -> float:
    r = 6371.0
    dlat, dlon = radians(lat2 - lat1), radians(lon2 - lon1)
    a = sin(dlat / 2) ** 2 + cos(radians(lat1)) * cos(radians(lat2)) * sin(dlon / 2) ** 2
    return 2 * r * asin(sqrt(a))


def hours_text(row) -> str:
    text = f"pon–sub {row['pon_sub_od']}–{row['pon_sub_do']}"
    if row["ned_od"]:
        text += f", ned {row['ned_od']}–{row['ned_do']}"
    return text


def route_url(lat0, lon0, lat1, lon1) -> str:
    return (
        "https://www.openstreetmap.org/directions?engine=fossgis_osrm_car"
        f"&route={lat0:.5f}%2C{lon0:.5f}%3B{lat1:.5f}%2C{lon1:.5f}"
    )


def is_content_dependent(item_hr: str) -> bool:
    name = item_hr.lower()
    return any(key in name for key in CONTENT_DEPENDENT)


def required_yard_flag(item_hr: str):
    name = item_hr.lower()
    for key, flag in YARD_FLAG_BY_KEYWORD.items():
        if key in name:
            return flag
    return None


def goes_to_yard(where_hr: str) -> bool:
    return "dvorišt" in (where_hr or "").lower()


@st.cache_data
def load_items():
    """Predmeti iz tablice: HR (B), EN (C), kategorija HR (D), gdje HR (F), napomena (I)."""
    from openpyxl import load_workbook

    ws = load_workbook(TABLE_PATH, data_only=True)["Predmeti"]
    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        hr, en = row[1], row[2]
        if hr and en:
            items.append(
                {
                    "hr": str(hr),
                    "en": str(en),
                    "kategorija": str(row[3] or ""),
                    "gdje": str(row[5] or ""),
                    "napomena": str(row[8] or ""),
                }
            )
    return items


@st.cache_resource(show_spinner="Učitavam model...")
def load_model():
    from transformers import CLIPModel, CLIPProcessor

    processor = CLIPProcessor.from_pretrained(MODEL_NAME)
    model = CLIPModel.from_pretrained(MODEL_NAME)
    model.eval()
    return processor, model


def classify(image, items):
    """Vraća [(indeks predmeta, vjerojatnost), ...] za TOP_N najvjerojatnijih."""
    import torch

    processor, model = load_model()
    prompts = [f"a photo of {it['en'].lower()}, a piece of waste" for it in items]
    with torch.no_grad():
        inputs = processor(text=prompts, images=image, return_tensors="pt", padding=True)
        probs = model(**inputs).logits_per_image.softmax(dim=-1)[0]
        top = torch.topk(probs, k=min(TOP_N, len(items)))
    return [(int(i), float(p)) for p, i in zip(top.values, top.indices)]


# ---------------------------------------------------------------- stranica: karta
def page_map(t, yards, now):
    st.info(t["click"])

    st.sidebar.header(t["filters"])
    types = st.sidebar.multiselect(
        t["type"],
        options=list(t["type_opts"].keys()),
        default=list(t["type_opts"].keys()),
        format_func=lambda k: t["type_opts"][k],
    )
    only_open = st.sidebar.checkbox(t["open_now"])
    extras = st.sidebar.multiselect(
        t["extra"],
        options=list(t["extra_opts"].keys()),
        format_func=lambda k: t["extra_opts"][k],
    )

    view = yards[yards["tip"].isin(types)]
    if only_open:
        view = view[view["otvoreno"]]
    for col in extras:
        view = view[view[col]]

    fmap = folium.Map(location=ZAGREB_CENTER, zoom_start=11, tiles="OpenStreetMap")
    for _, r in view.iterrows():
        status = t["open"] if r["otvoreno"] else t["closed"]
        popup = (
            f"<b>{r['naziv']}</b><br>{r['adresa']}<br>"
            f"{t['hours']}: {hours_text(r)}<br>"
            f"{t['phone']}: {r['telefon']}<br><b>{status}</b>"
        )
        folium.Marker(
            [r["lat"], r["lon"]],
            popup=folium.Popup(popup, max_width=300),
            tooltip=r["naziv"],
            icon=folium.Icon(
                color="green" if r["tip"] == "stalno" else "orange",
                icon="recycle",
                prefix="fa",
            ),
        ).add_to(fmap)

    if st.session_state.user_pos:
        folium.Marker(
            st.session_state.user_pos,
            tooltip=t["you"],
            icon=folium.Icon(color="blue", icon="user", prefix="fa"),
        ).add_to(fmap)

    result = st_folium(fmap, height=520, use_container_width=True, returned_objects=["last_clicked"])
    if result and result.get("last_clicked"):
        pos = (result["last_clicked"]["lat"], result["last_clicked"]["lng"])
        if pos != st.session_state.user_pos:
            st.session_state.user_pos = pos
            st.rerun()

    if view.empty:
        st.warning(t["none"])
    elif st.session_state.user_pos:
        lat0, lon0 = st.session_state.user_pos
        view = view.assign(
            km=view.apply(lambda r: haversine_km(lat0, lon0, r["lat"], r["lon"]), axis=1)
        ).sort_values("km")
        best = view.iloc[0]
        st.subheader(t["nearest"])
        c1, c2, c3 = st.columns(3)
        c1.metric(best["naziv"], t["open"] if best["otvoreno"] else t["closed"])
        c2.metric(t["distance"], f"{best['km']:.1f} km")
        c3.markdown(f"**{t['hours']}:** {hours_text(best)}  \n**{t['phone']}:** {best['telefon']}")
        st.link_button(t["route"], route_url(lat0, lon0, best["lat"], best["lon"]))

    st.caption(t["note"])


# ---------------------------------------------------------------- stranica: prepoznavanje
def show_nearest_yard(t, yards, flag):
    """Najbliže stalno dvorište (koje prima traženu vrstu otpada, ako je zadana)."""
    if not st.session_state.user_pos:
        st.info(t["no_pos"])
        return
    lat0, lon0 = st.session_state.user_pos
    pool = yards[yards["tip"] == "stalno"]
    if flag:
        pool = pool[pool[flag]]
    if pool.empty:
        st.warning(t["none"])
        return
    pool = pool.assign(
        km=pool.apply(lambda r: haversine_km(lat0, lon0, r["lat"], r["lon"]), axis=1)
    ).sort_values("km")
    best = pool.iloc[0]
    st.subheader(t["nearest"])
    c1, c2, c3 = st.columns(3)
    c1.metric(best["naziv"], t["open"] if best["otvoreno"] else t["closed"])
    c2.metric(t["distance"], f"{best['km']:.1f} km")
    c3.markdown(f"**{t['hours']}:** {hours_text(best)}  \n**{t['phone']}:** {best['telefon']}")
    st.link_button(t["route"], route_url(lat0, lon0, best["lat"], best["lon"]))


def page_scan(t, lang, yards):
    st.header(t["scan_title"])
    st.write(t["scan_intro"])

    try:
        from PIL import Image, ImageOps

        items = load_items()
        import torch  # noqa: F401
        import transformers  # noqa: F401
    except FileNotFoundError:
        st.error(t["no_table"])
        return
    except ImportError:
        st.warning(t["unavailable"])
        return

    source = st.radio(t["scan_title"], [t["src_cam"], t["src_upl"]], horizontal=True, label_visibility="collapsed")
    if source == t["src_cam"]:
        file = st.camera_input(t["take"])
    else:
        file = st.file_uploader(t["upload"], type=["jpg", "jpeg", "png", "webp"])
    if file is None:
        return

    image = ImageOps.exif_transpose(Image.open(file)).convert("RGB")
    st.image(image, width=300)

    with st.spinner(t["analysing"]):
        top = classify(image, items)

    name_key = "hr" if lang == "hr" else "en"
    best_idx, best_p = top[0]
    if best_p >= CONFIDENCE_THRESHOLD:
        st.success(f"{t['likely']}: **{items[best_idx][name_key]}** ({best_p:.0%})")
    else:
        st.warning(t["unsure"])

    options = [idx for idx, _ in top] + [-1]
    probs = dict(top)

    def label(idx):
        if idx == -1:
            return t["none_of"]
        return f"{items[idx][name_key]} ({probs[idx]:.0%})"

    choice = st.radio(t["pick"], options, format_func=label, index=0)
    if choice == -1:
        st.info(t["none_msg"])
        return

    item = items[choice]
    st.subheader(item[name_key])
    st.markdown(f"**{t['category']}:** {item['kategorija']}")
    st.markdown(f"**{t['where']}:** {item['gdje']}")
    if item["napomena"]:
        st.markdown(f"**{t['remark']}:** {item['napomena']}")

    if is_content_dependent(item["hr"]):
        st.warning(t["sensitive"])
        if not st.checkbox(t["confirm"]):
            return

    if goes_to_yard(item["gdje"]):
        show_nearest_yard(t, yards, required_yard_flag(item["hr"]))

    st.caption(t["disclaimer"])


# ---------------------------------------------------------------- glavni dio
def main():
    st.set_page_config(page_title="Zeleni rudnik", page_icon="♻️", layout="wide")

    if "user_pos" not in st.session_state:
        st.session_state.user_pos = None

    lang = st.sidebar.radio("Jezik / Language", ["hr", "en"], horizontal=True)
    t = TEXT[lang]

    st.title(t["title"])
    st.caption(t["subtitle"])

    yards = load_yards()
    now = datetime.now(ZoneInfo("Europe/Zagreb"))
    yards["otvoreno"] = yards.apply(lambda r: is_open(r, now), axis=1)

    page = st.sidebar.radio(t["page"], ["map", "scan"], format_func=lambda k: t["nav_map" if k == "map" else "nav_scan"])

    if page == "map":
        page_map(t, yards, now)
    else:
        page_scan(t, lang, yards)


if __name__ == "__main__":
    main()
